#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
zip_rar_folder2txt_split.py - نسخة مُحدَّثة تدعم TAR, GZ, PDF, Excel, Word
تُنشئ مجلداً وملفات منفصلة لكل ملف مستخرج
"""

import sys
import site

try:
    user_site = site.getusersitepackages()
    if user_site and user_site not in sys.path:
        sys.path.insert(0, user_site)
except Exception:
    pass

import os
import re
import json
import shutil
import hashlib
import zipfile
import tarfile
import gzip
import datetime
import tempfile
import sqlite3
import subprocess
import warnings
from pathlib import Path
from typing import Dict, List, Tuple, Optional, Union

warnings.filterwarnings("ignore", category=UserWarning)

# ============ استيراد المكتبات الاختيارية ============
def check_and_import(module_name, package_name=None):
    try:
        return __import__(module_name)
    except ImportError:
        if package_name is None:
            package_name = module_name
        print(f"⚠️ مكتبة {module_name} غير مثبتة. بعض الميزات لن تعمل.")
        print(f"   قم بتثبيتها باستخدام: pip install {package_name}")
        return None

pd = check_and_import('pandas', 'pandas')
if pd:
    try:
        from openpyxl import load_workbook
    except ImportError:
        pd = None
        print("⚠️ openpyxl غير مثبتة، معالجة Excel ستكون محدودة.")

docx = check_and_import('docx', 'python-docx')
bs4 = check_and_import('bs4', 'beautifulsoup4')
rarfile = check_and_import('rarfile', 'rarfile')
pdfplumber = check_and_import('pdfplumber', 'pdfplumber')
PIL = check_and_import('PIL', 'Pillow')
pytesseract = check_and_import('pytesseract', 'pytesseract')

# ============ الامتدادات المدعومة ============
TEXT_EXTENSIONS = {'.txt', '.py', '.js', '.json', '.xml', '.csv', '.md', '.yml', '.yaml', 
                   '.ini', '.cfg', '.conf', '.java', '.c', '.cpp', '.h', '.cs', '.php', 
                   '.rb', '.go', '.rs', '.swift', '.kt', '.sql', '.sh', '.bat', '.ps1',
                   '.r', '.m', '.f', '.for', '.f90', '.f95', '.properties', '.toml',
                   '.lock', '.log', '.tex', '.rst', '.adoc', '.asm', '.v', '.vhdl'}

DB_EXTENSIONS = {'.db', '.sqlite', '.sqlite3'}
WORD_EXTENSIONS = {'.docx'}
EXCEL_EXTENSIONS = {'.xls', '.xlsx'}
HTML_EXTENSIONS = {'.html', '.htm'}
PDF_EXTENSIONS = {'.pdf'}
ARCHIVE_EXTENSIONS = {'.zip', '.rar'}
TAR_EXTENSIONS = {'.tar', '.tar.gz', '.tgz', '.tar.bz2', '.tar.xz'}
GZ_EXTENSIONS = {'.gz'}

BINARY_EXTENSIONS = {'.pyc', '.pyo', '.pyd', '.so', '.dll', '.exe', '.bin', '.obj', 
                     '.o', '.a', '.lib', '.dylib', '.bundle', '.class', '.jar', 
                     '.war', '.ear', '.apk', '.ipa', '.app', '.dmg', '.iso', 
                     '.img', '.raw', '.dat', '.pkl', '.pickle', '.npy', '.npz',
                     '.pt', '.pth', '.h5', '.hdf', '.fits', '.parquet', '.feather',
                     '.msi', '.msu'}

# ============ دوال مساعدة ============
def safe_makedirs(path):
    if not os.path.exists(path):
        os.makedirs(path, exist_ok=True)

def get_unique_dirname(target_dir, base_name):
    dir_path = os.path.join(target_dir, base_name + "_extracted")
    if not os.path.exists(dir_path):
        return dir_path
    counter = 1
    while True:
        new_path = os.path.join(target_dir, f"{base_name}_extracted_{counter}")
        if not os.path.exists(new_path):
            return new_path
        counter += 1

def is_split_archive_extension(extension):
    ext_lower = extension.lower()
    if ext_lower.startswith('.z') and ext_lower[2:].isdigit():
        return True
    if ext_lower in ['.001', '.002', '.003', '.004', '.005', '.006', '.007', '.008', '.009', '.010',
                     '.r00', '.r01', '.r02', '.r03', '.r04', '.part1.rar', '.part2.rar', '.part3.rar']:
        return True
    return False

def should_ignore_file(file_path):
    parts = file_path.split('/')
    for part in parts:
        if part.startswith('.') and part != '.' and part != '..':
            return True
    if '__pycache__' in parts:
        return True
    if file_path.endswith('.pyc'):
        return True
    for i, part in enumerate(parts):
        if part == 'venv' and i < len(parts) - 1:
            return True
    if '/venv/' in file_path or file_path.startswith('venv/'):
        return True
    return False

def convert_timestamp(value):
    if pd is None:
        return str(value)
    try:
        if pd.isna(value) or value is None:
            return "NULL"
        if isinstance(value, (int, float)):
            try:
                value = datetime.datetime.fromtimestamp(value)
            except:
                pass
        if isinstance(value, str):
            formats = [
                '%Y-%m-%d %H:%M:%S',
                '%Y-%m-%d %H:%M:%S.%f',
                '%Y-%m-%dT%H:%M:%S',
                '%Y-%m-%dT%H:%M:%S.%f',
                '%Y-%m-%d',
                '%d/%m/%Y %H:%M:%S'
            ]
            for fmt in formats:
                try:
                    value = datetime.datetime.strptime(value, fmt)
                    break
                except:
                    continue
        if isinstance(value, datetime.datetime):
            hour = value.hour
            if hour == 0:
                hour_12 = 12
                meridiem = "ص"
            elif 1 <= hour <= 11:
                hour_12 = hour
                meridiem = "ص"
            elif hour == 12:
                hour_12 = 12
                meridiem = "م"
            else:
                hour_12 = hour - 12
                meridiem = "م"
            return f"{value.day:02d}/{value.month:02d}/{value.year} {hour_12}:{value.minute:02d}:{value.second:02d} {meridiem}"
        return str(value)
    except Exception:
        return str(value)

def extract_text_from_html(html_content):
    if bs4 is None:
        return html_content
    try:
        soup = bs4.BeautifulSoup(html_content, 'html.parser')
        for script in soup(["script", "style"]):
            script.decompose()
        text = soup.get_text()
        lines = (line.strip() for line in text.splitlines())
        chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
        text = '\n'.join(chunk for chunk in chunks if chunk)
        return text
    except Exception as e:
        return html_content

# ============ دوال معالجة قواعد البيانات ============
def export_db_to_excel(db_path, output_excel):
    if pd is None:
        print(" ❌ pandas غير مثبتة. لا يمكن التصدير إلى Excel.")
        return None, 0
    try:
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';")
        tables = [row[0] for row in cursor.fetchall()]
        cursor.close()
        if not tables:
            conn.close()
            return None, 0
        with pd.ExcelWriter(output_excel, engine='openpyxl') as writer:
            total_rows = 0
            for table in tables:
                try:
                    df = pd.read_sql_query(f"SELECT * FROM `{table}`", conn)
                    for col in df.columns:
                        df[col] = df[col].apply(convert_timestamp)
                    sheet_name = table[:31]
                    df.to_excel(writer, sheet_name=sheet_name, index=False)
                    total_rows += len(df)
                    print(f" ✓ جدول '{table}': {len(df)} صف")
                except Exception as e:
                    print(f" ⚠️ خطأ في تصدير '{table}': {str(e)}")
        conn.close()
        return output_excel, total_rows
    except Exception as e:
        print(f" ❌ خطأ في التصدير: {str(e)}")
        if 'conn' in locals():
            conn.close()
        return None, 0

def extract_db_via_excel_to_text(db_path, output_dir):
    print(f"🗄️→📊→📄 معالجة قاعدة البيانات عبر Excel: {os.path.basename(db_path)}")
    temp_excel = tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx').name
    try:
        excel_path, total_rows = export_db_to_excel(db_path, temp_excel)
        if not excel_path:
            return [], 0
        if pd is None:
            return [], 0
        excel_file = pd.ExcelFile(excel_path)
        sheet_names = excel_file.sheet_names
        files_created = []
        for sheet in sheet_names:
            df = pd.read_excel(excel_path, sheet_name=sheet)
            safe_sheet = re.sub(r'[\\/*?:"<>|]', '_', sheet)
            out_path = os.path.join(output_dir, f"{safe_sheet}.tsv")
            with open(out_path, 'w', encoding='utf-8') as f:
                f.write("\t".join(map(str, df.columns)) + "\n")
                for _, row in df.iterrows():
                    row_str = [str(val) if not pd.isna(val) else "NULL" for val in row]
                    f.write("\t".join(row_str) + "\n")
            files_created.append(out_path)
            print(f"   ✓ تم إنشاء {os.path.basename(out_path)}")
        return files_created, total_rows
    finally:
        try:
            if os.path.exists(temp_excel):
                os.unlink(temp_excel)
        except:
            pass

def extract_db_direct_to_text(db_path, output_dir):
    if not os.path.exists(db_path):
        return [], 0
    try:
        conn = sqlite3.connect(db_path)
        cursor = conn.cursor()
        cursor.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%';")
        tables = cursor.fetchall()
        if not tables:
            conn.close()
            return [], 0
        files_created = []
        total_rows = 0
        for table in tables:
            table_name = table[0]
            cursor.execute(f"SELECT * FROM \"{table_name}\"")
            rows = cursor.fetchall()
            col_names = [description[0] for description in cursor.description]
            safe_name = re.sub(r'[\\/*?:"<>|]', '_', table_name)
            out_path = os.path.join(output_dir, f"{safe_name}.tsv")
            with open(out_path, 'w', encoding='utf-8') as f:
                f.write("\t".join(col_names) + "\n")
                for row in rows:
                    row_str = []
                    for value in row:
                        if value is None:
                            row_str.append("NULL")
                        elif isinstance(value, (bytes, bytearray)):
                            row_str.append("<binary data>")
                        else:
                            row_str.append(str(value))
                    f.write("\t".join(row_str) + "\n")
            files_created.append(out_path)
            total_rows += len(rows)
            print(f" ✓ جدول {table_name}: {len(rows)} صف")
        conn.close()
        return files_created, total_rows
    except sqlite3.Error as e:
        print(f" ❌ خطأ في قاعدة البيانات: {str(e)}")
        return [], 0

# ============ دوال معالجة Excel ============
def extract_excel_to_text(excel_path, output_dir):
    if pd is None:
        print(" ❌ pandas غير مثبتة. لا يمكن معالجة Excel.")
        return [], 0
    if not os.path.exists(excel_path):
        return [], 0
    try:
        excel_file = pd.ExcelFile(excel_path)
        sheet_names = excel_file.sheet_names
        files_created = []
        total_rows = 0
        for sheet in sheet_names:
            df = pd.read_excel(excel_path, sheet_name=sheet)
            safe_sheet = re.sub(r'[\\/*?:"<>|]', '_', sheet)
            out_path = os.path.join(output_dir, f"{safe_sheet}.tsv")
            with open(out_path, 'w', encoding='utf-8') as f:
                f.write("\t".join(map(str, df.columns)) + "\n")
                for _, row in df.iterrows():
                    row_str = [str(val) if not pd.isna(val) else "NULL" for val in row]
                    f.write("\t".join(row_str) + "\n")
            files_created.append(out_path)
            total_rows += len(df)
            print(f" ✓ ورقة {sheet}: {len(df)} صف")
        return files_created, total_rows
    except Exception as e:
        print(f" ❌ خطأ في Excel: {str(e)}")
        return [], 0

# ============ دوال معالجة Word ============
def extract_docx_to_text(docx_path, output_dir):
    if docx is None:
        print(" ❌ python-docx غير مثبتة. لا يمكن معالجة Word.")
        return [], 0
    if not os.path.exists(docx_path):
        return [], 0
    try:
        doc = docx.Document(docx_path)
        paragraphs = []
        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append(para.text)
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                tables_text.append("\t".join(cells))
        out_path = os.path.join(output_dir, "document_text.txt")
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write("\n".join(paragraphs))
            if tables_text:
                f.write("\n\n--- جداول ---\n")
                f.write("\n".join(tables_text))
        return [out_path], len(paragraphs) + len(tables_text)
    except Exception as e:
        print(f" ❌ خطأ في Word: {str(e)}")
        return [], 0

# ============ دوال معالجة HTML ============
def extract_html_to_text(html_path, output_dir):
    if bs4 is None:
        try:
            with open(html_path, 'r', encoding='utf-8') as f:
                content = f.read()
            text = re.sub(r'<[^>]+>', '', content)
            text = re.sub(r'\s+', ' ', text).strip()
            out_path = os.path.join(output_dir, "html_text.txt")
            with open(out_path, 'w', encoding='utf-8') as f:
                f.write(text)
            return [out_path], 1
        except Exception as e:
            return [], 0
    try:
        with open(html_path, 'r', encoding='utf-8') as f:
            soup = bs4.BeautifulSoup(f, 'html.parser')
        for tag in soup(['script', 'style', 'head', 'title', 'meta', '[document]']):
            tag.decompose()
        text = soup.get_text(separator='\n', strip=True)
        out_path = os.path.join(output_dir, "html_text.txt")
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(text)
        lines = len(text.split('\n'))
        return [out_path], lines
    except Exception as e:
        print(f" ❌ خطأ في HTML: {str(e)}")
        return [], 0

# ============ دوال معالجة PDF المتقدمة ============
def extract_pdf_advanced(pdf_path, output_dir, use_ocr=True):
    if pdfplumber is None:
        print(" ❌ pdfplumber غير مثبتة. لا يمكن معالجة PDF.")
        return [], 0, 0
    
    base_name = os.path.splitext(os.path.basename(pdf_path))[0]
    pdf_output_dir = os.path.join(output_dir, f"{base_name}_pdf_extracted")
    safe_makedirs(pdf_output_dir)
    
    images_dir = os.path.join(pdf_output_dir, "images")
    tables_dir = os.path.join(pdf_output_dir, "tables")
    safe_makedirs(images_dir)
    safe_makedirs(tables_dir)
    
    full_text_path = os.path.join(pdf_output_dir, "full_text.txt")
    summary_path = os.path.join(pdf_output_dir, "summary.json")
    
    try:
        with pdfplumber.open(pdf_path) as pdf:
            total_pages = len(pdf.pages)
            pages_with_ocr = 0
            total_images = 0
            total_tables = 0
            metadata = pdf.metadata or {}
            
            with open(full_text_path, 'w', encoding='utf-8') as txt_out:
                txt_out.write("=" * 80 + "\n")
                txt_out.write(f"محتوى PDF: {os.path.basename(pdf_path)}\n")
                txt_out.write(f"عدد الصفحات: {total_pages}\n")
                txt_out.write("=" * 80 + "\n\n")
                
                for page_num, page in enumerate(pdf.pages, 1):
                    txt_out.write(f"\n{'='*80}\n")
                    txt_out.write(f"الصفحة {page_num}\n")
                    txt_out.write(f"{'='*80}\n\n")
                    
                    page_text = page.extract_text()
                    if page_text and page_text.strip():
                        txt_out.write(page_text)
                        txt_out.write("\n\n")
                    else:
                        txt_out.write("[الصفحة تحتوي على صور، قد تحتاج OCR]\n")
                        pages_with_ocr += 1
                    
                    tables = page.extract_tables()
                    for i, table in enumerate(tables):
                        if table and any(any(row) for row in table):
                            total_tables += 1
                            csv_path = os.path.join(tables_dir, f"page{page_num:04d}_table{i+1:02d}.csv")
                            with open(csv_path, 'w', encoding='utf-8') as f:
                                for row in table:
                                    f.write("\t".join([str(cell) if cell else "" for cell in row]) + "\n")
                            md_path = os.path.join(tables_dir, f"page{page_num:04d}_table{i+1:02d}.md")
                            with open(md_path, 'w', encoding='utf-8') as f:
                                f.write(f"# جدول من الصفحة {page_num}\n\n")
                                if table and table[0]:
                                    f.write("| " + " | ".join([str(h) if h else "" for h in table[0]]) + " |\n")
                                    f.write("|" + "|".join(["---"] * len(table[0])) + "|\n")
                                    for row in table[1:]:
                                        f.write("| " + " | ".join([str(cell) if cell else "" for cell in row]) + " |\n")
                    
                    if hasattr(page, 'images') and page.images:
                        for img_idx, img in enumerate(page.images):
                            total_images += 1
                            img_info = {
                                'page': page_num,
                                'index': img_idx + 1,
                                'bbox': img.get('bbox'),
                                'width': img.get('width'),
                                'height': img.get('height'),
                                'alt_text': img.get('alt', '')
                            }
                            img_json_path = os.path.join(images_dir, f"page{page_num:04d}_img{img_idx+1:02d}.json")
                            with open(img_json_path, 'w', encoding='utf-8') as f:
                                json.dump(img_info, f, ensure_ascii=False, indent=2)
            
            if use_ocr and pytesseract and PIL:
                print("   🖼️ تشغيل OCR على الصفحات...")
                for page_num, page in enumerate(pdf.pages, 1):
                    if not page.extract_text() or not page.extract_text().strip():
                        try:
                            im = page.to_image(resolution=300)
                            img_path = os.path.join(images_dir, f"page{page_num:04d}.png")
                            im.save(img_path, format="PNG")
                            ocr_text = pytesseract.image_to_string(PIL.Image.open(img_path), lang='ara+eng')
                            with open(full_text_path, 'a', encoding='utf-8') as f:
                                f.write(f"\n[OCR للصفحة {page_num}]\n")
                                f.write(ocr_text + "\n")
                            pages_with_ocr += 1
                            total_images += 1
                        except Exception as e:
                            print(f"      ⚠️ فشل OCR للصفحة {page_num}: {e}")
            
            summary = {
                'pdf_file': os.path.basename(pdf_path),
                'total_pages': total_pages,
                'pages_with_ocr': pages_with_ocr,
                'total_images': total_images,
                'total_tables': total_tables,
                'metadata': metadata,
                'output_dir': pdf_output_dir
            }
            with open(summary_path, 'w', encoding='utf-8') as f:
                json.dump(summary, f, ensure_ascii=False, indent=2)
            
            print(f"   ✅ PDF معالج: {os.path.basename(pdf_path)}")
            print(f"      الصفحات: {total_pages}, صور: {total_images}, جداول: {total_tables}")
            return [full_text_path, summary_path], total_pages, 0
        
    except Exception as e:
        print(f" ❌ خطأ في معالجة PDF: {str(e)}")
        return [], 0, 1

# ============ دوال معالجة الأرشيفات ============
def extract_archive_to_files(archive_path, output_dir, archive_type="zip"):
    if not os.path.exists(archive_path):
        return [], 0, 0
    try:
        if archive_type == "zip":
            if not zipfile.is_zipfile(archive_path):
                return [], 0, 0
            archive = zipfile.ZipFile(archive_path, 'r')
        elif archive_type == "rar":
            if rarfile is None:
                print(" ❌ rarfile غير مثبتة. لا يمكن معالجة RAR.")
                return [], 0, 0
            try:
                archive = rarfile.RarFile(archive_path, 'r')
            except:
                return [], 0, 0
        else:
            return [], 0, 0
        
        file_list = archive.namelist()
        files_processed = 0
        files_skipped = 0
        created_files = []
        
        for file_name in sorted(file_list):
            if file_name.endswith('/'):
                continue
            if should_ignore_file(file_name):
                files_skipped += 1
                continue
            ext = pathlib.Path(file_name).suffix.lower()
            if ext in EXCEL_EXTENSIONS or ext in WORD_EXTENSIONS or ext in PDF_EXTENSIONS:
                files_skipped += 1
                continue
            if ext in BINARY_EXTENSIONS:
                files_skipped += 1
                continue
            try:
                if archive_type == "zip":
                    content_bytes = archive.read(file_name)
                else:
                    with archive.open(file_name, 'r') as f:
                        content_bytes = f.read()
                if len(content_bytes) == 0:
                    continue
                try:
                    content = content_bytes.decode('utf-8')
                except UnicodeDecodeError:
                    try:
                        content = content_bytes.decode('latin-1')
                    except:
                        files_skipped += 1
                        continue
                dest_path = os.path.join(output_dir, file_name)
                dest_dir = os.path.dirname(dest_path)
                safe_makedirs(dest_dir)
                with open(dest_path, 'w', encoding='utf-8') as f:
                    f.write(content)
                created_files.append(dest_path)
                files_processed += 1
            except Exception as e:
                print(f" ⚠️ خطأ في استخراج {file_name}: {str(e)}")
                files_skipped += 1
        return created_files, files_processed, files_skipped
    except Exception as e:
        print(f" ❌ حدث خطأ في الأرشيف: {str(e)}")
        return [], 0, 0

def extract_tar_to_files(tar_path, output_dir):
    if not os.path.exists(tar_path):
        return [], 0, 0
    
    ext = pathlib.Path(tar_path).suffix.lower()
    if ext == '.tar':
        mode = 'r'
    elif ext == '.gz' or ext == '.tgz' or tar_path.endswith('.tar.gz'):
        mode = 'r:gz'
    elif ext == '.bz2' or tar_path.endswith('.tar.bz2'):
        mode = 'r:bz2'
    elif ext == '.xz' or tar_path.endswith('.tar.xz'):
        mode = 'r:xz'
    else:
        mode = 'r:*'
    
    try:
        with tarfile.open(tar_path, mode) as tar:
            members = tar.getmembers()
            processed = 0
            skipped = 0
            created_files = []
            
            for member in members:
                if member.isfile():
                    if should_ignore_file(member.name):
                        skipped += 1
                        continue
                    file_ext = pathlib.Path(member.name).suffix.lower()
                    if file_ext in EXCEL_EXTENSIONS or file_ext in WORD_EXTENSIONS or file_ext in PDF_EXTENSIONS:
                        skipped += 1
                        continue
                    if file_ext in BINARY_EXTENSIONS:
                        skipped += 1
                        continue
                    try:
                        f = tar.extractfile(member)
                        if f:
                            content_bytes = f.read()
                            f.close()
                            if len(content_bytes) == 0:
                                continue
                            try:
                                content = content_bytes.decode('utf-8')
                            except UnicodeDecodeError:
                                try:
                                    content = content_bytes.decode('latin-1')
                                except:
                                    skipped += 1
                                    continue
                            dest_path = os.path.join(output_dir, member.name)
                            dest_dir = os.path.dirname(dest_path)
                            safe_makedirs(dest_dir)
                            with open(dest_path, 'w', encoding='utf-8') as f_out:
                                f_out.write(content)
                            created_files.append(dest_path)
                            processed += 1
                    except Exception as e:
                        print(f" ⚠️ خطأ في استخراج {member.name}: {str(e)}")
                        skipped += 1
            return created_files, processed, skipped
    except Exception as e:
        print(f" ❌ خطأ في معالجة TAR: {str(e)}")
        return [], 0, 0

def extract_gz_to_file(gz_path, output_dir):
    try:
        with gzip.open(gz_path, 'rt', encoding='utf-8', errors='ignore') as f:
            content = f.read()
        if len(content.strip()) == 0:
            return [], 0, 1
        out_filename = os.path.splitext(os.path.basename(gz_path))[0] + ".txt"
        out_path = os.path.join(output_dir, out_filename)
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(content)
        return [out_path], 1, 0
    except Exception as e:
        print(f" ❌ خطأ في معالجة GZ: {str(e)}")
        return [], 0, 1

# ============ معالجة ملف واحد عادي (نصي) ============
def extract_single_file_to_text(file_path, output_dir):
    if not os.path.isfile(file_path):
        return [], 0, 0
    ext = pathlib.Path(file_path).suffix.lower()
    if ext in BINARY_EXTENSIONS and ext not in DB_EXTENSIONS and ext not in WORD_EXTENSIONS and ext not in EXCEL_EXTENSIONS and ext not in HTML_EXTENSIONS and ext not in PDF_EXTENSIONS:
        return [], 0, 1
    try:
        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()
        if len(content.strip()) == 0:
            return [], 0, 1
        out_path = os.path.join(output_dir, os.path.basename(file_path) + ".txt")
        with open(out_path, 'w', encoding='utf-8') as f:
            f.write(content)
        return [out_path], 1, 0
    except Exception as e:
        print(f" ❌ خطأ في معالجة {file_path}: {str(e)}")
        return [], 0, 1

# ============ المعالج الرئيسي ============
def process_single_item(item_path, via_excel=False, use_ocr=False):
    results = []
    if not os.path.exists(item_path):
        print(f"❌ المسار غير موجود: {item_path}")
        return results
    
    output_dir = os.path.dirname(item_path)
    if not output_dir:
        output_dir = '/'
    base_name = os.path.basename(item_path)
    name_without_ext = os.path.splitext(base_name)[0]
    
    target_dir = get_unique_dirname(output_dir, name_without_ext)
    safe_makedirs(target_dir)
    print(f"📁 سيتم حفظ المخرجات في: {target_dir}")
    
    if os.path.isfile(item_path):
        file_ext = pathlib.Path(item_path).suffix.lower()
        
        if is_split_archive_extension(file_ext):
            print(f"⚠️ ملف جزء من أرشيف متعدد: {base_name} - سيتم تجاهله")
            return results
        
        if file_ext in DB_EXTENSIONS:
            if via_excel:
                print(f"🗄️→📊→📄 معالجة قاعدة بيانات عبر Excel: {base_name}")
                files, total_rows = extract_db_via_excel_to_text(item_path, target_dir)
                results.append((files, total_rows, 0))
            else:
                print(f"🗄️ معالجة قاعدة بيانات مباشرة: {base_name}")
                files, total_rows = extract_db_direct_to_text(item_path, target_dir)
                results.append((files, total_rows, 0))
        
        elif file_ext in EXCEL_EXTENSIONS:
            print(f"📊 معالجة ملف Excel: {base_name}")
            files, total_rows = extract_excel_to_text(item_path, target_dir)
            results.append((files, total_rows, 0))
        
        elif file_ext in WORD_EXTENSIONS:
            print(f"📝 معالجة مستند Word: {base_name}")
            files, total_rows = extract_docx_to_text(item_path, target_dir)
            results.append((files, total_rows, 0))
        
        elif file_ext in HTML_EXTENSIONS:
            print(f"🌐 استخراج نص من HTML: {base_name}")
            files, lines = extract_html_to_text(item_path, target_dir)
            results.append((files, lines, 0))
        
        elif file_ext in PDF_EXTENSIONS:
            print(f"📄 معالجة PDF (متقدمة): {base_name}")
            files, pages, err = extract_pdf_advanced(item_path, target_dir, use_ocr=use_ocr)
            results.append((files, pages, err))
        
        elif file_ext in TAR_EXTENSIONS or item_path.endswith('.tar.gz'):
            print(f"📦 معالجة ملف TAR: {base_name}")
            files, processed, skipped = extract_tar_to_files(item_path, target_dir)
            results.append((files, processed, skipped))
        
        elif file_ext == '.gz' and not item_path.endswith('.tar.gz'):
            print(f"📦 معالجة ملف GZ: {base_name}")
            files, processed, skipped = extract_gz_to_file(item_path, target_dir)
            results.append((files, processed, skipped))
        
        elif file_ext == '.zip':
            print(f"📦 معالجة ملف ZIP: {base_name}")
            files, processed, skipped = extract_archive_to_files(item_path, target_dir, "zip")
            results.append((files, processed, skipped))
        elif file_ext == '.rar' and rarfile is not None:
            print(f"📦 معالجة ملف RAR: {base_name}")
            files, processed, skipped = extract_archive_to_files(item_path, target_dir, "rar")
            results.append((files, processed, skipped))
        elif file_ext == '.rar' and rarfile is None:
            print(f"⚠️ ملف RAR يتجاهل (rarfile غير مثبت)")
            results.append(([], 0, 1))
        
        else:
            if file_ext in TEXT_EXTENSIONS or file_ext not in BINARY_EXTENSIONS:
                print(f"📄 معالجة ملف نصي: {base_name}")
                files, processed, skipped = extract_single_file_to_text(item_path, target_dir)
                results.append((files, processed, skipped))
            else:
                print(f"❌ نوع الملف غير مدعوم أو ثنائي: {base_name}")
    
    elif os.path.isdir(item_path):
        print(f"📁 معالجة المجلد: {base_name}")
        all_files = []
        processed = 0
        skipped = 0
        for root, dirs, files in os.walk(item_path):
            dirs[:] = [d for d in dirs if not d.startswith('.') and d != '__pycache__' and d != 'venv']
            for file in files:
                full_path = os.path.join(root, file)
                rel_path = os.path.relpath(full_path, item_path)
                if should_ignore_file(rel_path):
                    skipped += 1
                    continue
                ext = pathlib.Path(file).suffix.lower()
                try:
                    if ext in DB_EXTENSIONS:
                        f, cnt = extract_db_direct_to_text(full_path, target_dir)
                        all_files.extend(f)
                        processed += cnt
                    elif ext in EXCEL_EXTENSIONS:
                        f, cnt = extract_excel_to_text(full_path, target_dir)
                        all_files.extend(f)
                        processed += cnt
                    elif ext in WORD_EXTENSIONS:
                        f, cnt = extract_docx_to_text(full_path, target_dir)
                        all_files.extend(f)
                        processed += cnt
                    elif ext in HTML_EXTENSIONS:
                        f, cnt = extract_html_to_text(full_path, target_dir)
                        all_files.extend(f)
                        processed += cnt
                    elif ext in PDF_EXTENSIONS:
                        f, cnt, err = extract_pdf_advanced(full_path, target_dir, use_ocr=use_ocr)
                        all_files.extend(f)
                        processed += cnt
                    elif ext in TAR_EXTENSIONS or full_path.endswith('.tar.gz'):
                        f, p, s = extract_tar_to_files(full_path, target_dir)
                        all_files.extend(f)
                        processed += p
                        skipped += s
                    elif ext == '.gz' and not full_path.endswith('.tar.gz'):
                        f, p, s = extract_gz_to_file(full_path, target_dir)
                        all_files.extend(f)
                        processed += p
                        skipped += s
                    elif ext == '.zip':
                        f, p, s = extract_archive_to_files(full_path, target_dir, "zip")
                        all_files.extend(f)
                        processed += p
                        skipped += s
                    elif ext == '.rar' and rarfile is not None:
                        f, p, s = extract_archive_to_files(full_path, target_dir, "rar")
                        all_files.extend(f)
                        processed += p
                        skipped += s
                    else:
                        with open(full_path, 'r', encoding='utf-8', errors='ignore') as f:
                            content = f.read()
                        if content.strip():
                            dest_path = os.path.join(target_dir, rel_path + ".txt")
                            dest_dir = os.path.dirname(dest_path)
                            safe_makedirs(dest_dir)
                            with open(dest_path, 'w', encoding='utf-8') as f_out:
                                f_out.write(content)
                            all_files.append(dest_path)
                            processed += 1
                        else:
                            skipped += 1
                except Exception as e:
                    print(f" ⚠️ خطأ في معالجة {rel_path}: {str(e)}")
                    skipped += 1
        results.append((all_files, processed, skipped))
    
    else:
        print(f"❌ نوع غير معروف: {item_path}")
    
    return results

def main():
    print("🔍 فحص المكتبات المثبتة:")
    print(f" ✓ zipfile: مثبت")
    print(f" ✓ tarfile: مثبت")
    print(f" ✓ gzip: مثبت")
    print(f" ✓ sqlite3: مثبت")
    print(f" ✓ python-docx: {'مثبت' if docx else 'غير مثبت'}")
    print(f" ✓ pandas: {'مثبت' if pd else 'غير مثبت'}")
    print(f" ✓ beautifulsoup4: {'مثبت' if bs4 else 'غير مثبت'}")
    print(f" ✓ rarfile: {'مثبت' if rarfile else 'غير مثبت'}")
    print(f" ✓ pdfplumber: {'مثبت' if pdfplumber else 'غير مثبت'}")
    print(f" ✓ PIL: {'مثبت' if PIL else 'غير مثبت'}")
    print(f" ✓ pytesseract: {'مثبت' if pytesseract else 'غير مثبت (OCR معطل)'}")
    print()
    
    if len(sys.argv) < 2:
        print("=" * 60)
        print("الاستخدام:")
        print("   python script.py [--via-excel] [--ocr] ملف1 ملف2 ...")
        print("الخيارات:")
        print("   --via-excel : تحويل قواعد البيانات عبر Excel")
        print("   --ocr       : تشغيل OCR على صفحات PDF التي لا تحتوي على نص")
        print("=" * 60)
        input("اضغط Enter للخروج...")
        return
    
    via_excel = "--via-excel" in sys.argv
    use_ocr = "--ocr" in sys.argv
    args = [arg for arg in sys.argv[1:] if arg not in ("--via-excel", "--ocr")]
    
    if via_excel:
        print("💡 استخدام التحويل عبر Excel للقواعد البيانات")
    if use_ocr:
        print("💡 تشغيل OCR على PDF")
    
    all_files_created = []
    total_processed = 0
    total_skipped = 0
    
    print(f"\n🎯 تم العثور على {len(args)} عنصر للمعالجة:")
    for i, item in enumerate(args, 1):
        print(f"\n[{i}/{len(args)}] {'='*50}")
        results = process_single_item(item, via_excel=via_excel, use_ocr=use_ocr)
        for files, proc, skip in results:
            all_files_created.extend(files)
            total_processed += proc
            total_skipped += skip
        print(f"✓ اكتمل: {len(files)} ملف منشأ")
    
    print("\n" + "=" * 60)
    print("📊 ملخص المعالجة النهائي:")
    print("=" * 60)
    print(f"📁 عدد الملفات النصية المنشأة: {len(all_files_created)}")
    print(f"📄 إجمالي العناصر المعالجة: {total_processed}")
    print(f"🚫 إجمالي العناصر المتجاهلة: {total_skipped}")
    
    print("\n✅ اكتملت المعالجة!")
    input("اضغط Enter للخروج...")

if __name__ == "__main__":
    main()
